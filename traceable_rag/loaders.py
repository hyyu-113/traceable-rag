import logging
import re
from abc import ABC, abstractmethod
from collections.abc import Iterator
from pathlib import Path
from zipfile import ZipFile

from docx import Document as WordDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pdfminer.converter import PDFPageAggregator
from pdfminer.layout import LAParams, LTContainer, LTTextBox
from pdfminer.pdfdocument import PDFDocument, PDFPasswordIncorrect
from pdfminer.pdfinterp import PDFPageInterpreter, PDFResourceManager
from pdfminer.pdfpage import PDFPage
from pdfminer.pdfparser import PDFParser

from traceable_rag.domain import Document, DocumentBlock, SourceLocation
from traceable_rag.utils import DocumentParseError, UnsupportedFileTypeError, sha256

logger = logging.getLogger(__name__)


class BaseLoader(ABC):
    @abstractmethod
    def load(self, document: Document) -> list[DocumentBlock]:
        """返回带原始定位的结构块。"""

    def block(self, document: Document, kind: str, text: str, location: SourceLocation) -> DocumentBlock:
        identity = f"{document.document_id}:{location.model_dump_json()}:{text}"
        return DocumentBlock(
            block_id=sha256(identity.encode()),
            document_id=document.document_id,
            block_type=kind,
            text=text.strip(),
            source_location=location,
            metadata={"file_name": document.file_name},
        )


class PdfLoader(BaseLoader):
    """提取文字层和块坐标，不执行 OCR 或表格结构识别。"""

    @staticmethod
    def _text_boxes(container: LTContainer) -> Iterator[LTTextBox]:
        # Form XObject 中也可能包含正文；只取文本框，不重复收集其内部文本行。
        for item in container:
            if isinstance(item, LTTextBox):
                yield item
            elif isinstance(item, LTContainer):
                yield from PdfLoader._text_boxes(item)

    def load(self, document: Document) -> list[DocumentBlock]:
        blocks = []
        try:
            with open(document.file_path, "rb") as source:
                pdf = PDFDocument(PDFParser(source))
                if pdf.encryption:
                    raise DocumentParseError("不支持加密 PDF，请先解除密码。")
                resources = PDFResourceManager(caching=False)
                device = PDFPageAggregator(resources, laparams=LAParams(all_texts=True))
                try:
                    interpreter = PDFPageInterpreter(resources, device)
                    for page_number, page in enumerate(PDFPage.create_pages(pdf), 1):
                        interpreter.process_page(page)
                        layout = device.get_result()
                        boxes = sorted(self._text_boxes(layout), key=lambda item: (-item.y1, item.x0))
                        for item in boxes:
                            text = item.get_text().strip()
                            if not text:
                                continue
                            # PDFMiner 以左下角为原点；统一转换成页面左上角坐标，保持接口约定。
                            bbox = (item.x0 - layout.x0, layout.y1 - item.y1, item.x1 - layout.x0, layout.y1 - item.y0)
                            location = SourceLocation(page_number=page_number, bbox=bbox)
                            block = self.block(document, "paragraph", text, location)
                            block.metadata.update({"page_width": layout.width, "page_height": layout.height, "pdf_parser": "pdfminer.six"})
                            blocks.append(block)
                finally:
                    device.close()
        except PDFPasswordIncorrect as exc:
            raise DocumentParseError("不支持加密 PDF，请先解除密码。") from exc
        return blocks


class DocxLoader(BaseLoader):
    def load(self, document: Document) -> list[DocumentBlock]:
        word = WordDocument(document.file_path)
        blocks = []
        headings: dict[int, str] = {}
        paragraph_index, table_index = 0, 0
        # 按正文 XML 顺序遍历，避免先提取全部段落导致表格章节错位。
        for element in word.iter_inner_content():
            section = " / ".join(headings.values()) or None
            if isinstance(element, Paragraph):
                paragraph_index += 1
                text = element.text.strip()
                if not text:
                    continue
                match = re.search(r"(?:Heading|标题)\s*(\d+)", element.style.name, re.IGNORECASE)
                if match:
                    level = int(match.group(1))
                    headings = {key: value for key, value in headings.items() if key < level}
                    headings[level] = text
                    section = " / ".join(headings.values())
                location = SourceLocation(section_path=section, paragraph_index=paragraph_index)
                blocks.append(self.block(document, "heading" if match else "paragraph", text, location))
            elif isinstance(element, Table):
                table_index += 1
                # 小批次保留整行；重复表头的原始行号单独记录。
                rows = [[cell.text.replace("\n", "；") for cell in row.cells] for row in element.rows]
                for start in range(0, len(rows), 15):
                    batch = rows[start : start + 15]
                    location = SourceLocation(
                        section_path=section,
                        table_index=table_index,
                        row_start=start + 1,
                        row_end=start + len(batch),
                        column_start=1,
                        column_end=max(map(len, batch)),
                    )
                    block = self.block(document, "table", "\n".join(" | ".join(row) for row in batch), location)
                    if start:
                        block.metadata["table_header"] = " | ".join(rows[0])
                        block.metadata["header_row"] = 1
                    blocks.append(block)
        return blocks


class ExcelLoader(BaseLoader):
    def load(self, document: Document) -> list[DocumentBlock]:
        workbook = load_workbook(document.file_path, read_only=True, data_only=False)
        blocks = []
        try:
            for sheet in workbook:
                if sheet.max_row * sheet.max_column > 2_000_000:
                    raise DocumentParseError("工作表有效范围过大，请清除多余格式后上传。")
                pending: list[tuple[int, list[str]]] = []
                for number, values in enumerate(sheet.iter_rows(values_only=True), 1):
                    row = ["" if value is None else str(value) for value in values]
                    if any(row):
                        pending.append((number, row))
                    if pending and (not any(row) or len(pending) >= 15):
                        blocks.append(self._region(document, sheet.title, pending))
                        pending = []
                if pending:
                    blocks.append(self._region(document, sheet.title, pending))
        finally:
            workbook.close()
        return blocks

    def _region(self, document: Document, sheet: str, rows: list[tuple[int, list[str]]]) -> DocumentBlock:
        columns = [i for _, row in rows for i, value in enumerate(row) if value]
        left, right = min(columns), max(columns)
        cell_range = f"{get_column_letter(left + 1)}{rows[0][0]}:{get_column_letter(right + 1)}{rows[-1][0]}"
        text = "\n".join(" | ".join(row[left : right + 1]) for _, row in rows)
        location = SourceLocation(
            sheet_name=sheet, cell_range=cell_range, row_start=rows[0][0], row_end=rows[-1][0], column_start=left + 1, column_end=right + 1
        )
        return self.block(document, "table", text, location)


class LoaderFactory:
    @staticmethod
    def create(extension: str) -> BaseLoader:
        loaders = {".pdf": PdfLoader, ".docx": DocxLoader, ".xlsx": ExcelLoader}
        if extension.lower() not in loaders:
            raise UnsupportedFileTypeError("仅支持 PDF、DOCX、XLSX 文件。")
        return loaders[extension.lower()]()

    @staticmethod
    def load(document: Document) -> list[DocumentBlock]:
        loader = LoaderFactory.create(Path(document.file_path).suffix)
        try:
            if document.file_type in ("docx", "xlsx"):
                with ZipFile(document.file_path) as archive:
                    if sum(item.file_size for item in archive.infolist()) > 100_000_000:
                        raise DocumentParseError("文档解压后过大。")
            blocks = loader.load(document)
            if not blocks:
                raise DocumentParseError("未提取到文字。扫描 PDF 需要 OCR，当前版本不支持。")
            logger.info("解析完成 document=%s blocks=%d", document.document_id, len(blocks))
            return blocks
        except DocumentParseError:
            raise
        except Exception as exc:
            raise DocumentParseError("文档解析失败，请检查文件是否完整。") from exc
