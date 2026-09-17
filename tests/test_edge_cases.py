from pathlib import Path

import pytest
from openpyxl import Workbook
from reportlab.pdfgen.canvas import Canvas
from test_loaders import document

from traceable_rag.loaders import LoaderFactory
from traceable_rag.storage import MetadataStore
from traceable_rag.utils import DocumentParseError


def test_blank_pdf_refuses(tmp_path: Path) -> None:
    path = tmp_path / "blank.pdf"
    canvas = Canvas(str(path))
    canvas.showPage()
    canvas.save()
    with pytest.raises(DocumentParseError, match="OCR"):
        LoaderFactory.load(document(path))


def test_excel_gaps_and_batches(tmp_path: Path) -> None:
    path = tmp_path / "regions.xlsx"
    book = Workbook()
    sheet = book.active
    for row in range(3, 23):
        sheet.cell(row, 2, f"产品{row}")
        sheet.cell(row, 3, row)
    sheet.cell(25, 2, "另一块")
    book.save(path)
    blocks = LoaderFactory.load(document(path))
    assert [item.source_location.cell_range for item in blocks] == ["B3:C17", "B18:C22", "B25:B25"]


def test_embedding_profile_change_blocked(tmp_path: Path) -> None:
    store = MetadataStore(tmp_path / "db")
    store.verify_embedding_profile("model-a")
    store.verify_embedding_profile("model-a")
    with pytest.raises(ValueError, match="Embedding"):
        store.verify_embedding_profile("model-b")
