from pathlib import Path

from docx import Document as WordDocument
from openpyxl import Workbook
from reportlab.pdfgen.canvas import Canvas

from traceable_rag.domain import Document
from traceable_rag.loaders import LoaderFactory
from traceable_rag.utils import sha256


def document(path: Path) -> Document:
    digest = sha256(path.read_bytes())
    return Document(document_id=digest, file_hash=digest, file_name=path.name, file_type=path.suffix[1:], file_path=str(path))


def test_pdf(tmp_path: Path) -> None:
    path = tmp_path / "sample.pdf"
    canvas = Canvas(str(path), pagesize=(600, 800))
    canvas.drawString(72, 728, "Travel policy")
    canvas.drawString(72, 630, "Hotel: 600")
    canvas.save()
    blocks = LoaderFactory.load(document(path))
    assert len(blocks) == 2
    assert blocks[0].source_location.page_number == 1
    assert blocks[0].source_location.bbox is not None


def test_docx_order(tmp_path: Path) -> None:
    path = tmp_path / "policy.docx"
    word = WordDocument()
    word.add_heading("差旅制度", 1)
    word.add_paragraph("住宿标准为600元")
    table = word.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "北京", "600"
    word.add_heading("采购制度", 1)
    word.save(path)
    blocks = LoaderFactory.load(document(path))
    assert blocks[2].source_location.section_path == "差旅制度"
    assert blocks[2].source_location.table_index == 1
    assert "600" in blocks[2].text
    assert blocks[1].source_location.paragraph_index == 2


def test_excel_region(tmp_path: Path) -> None:
    path = tmp_path / "prices.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "价格"
    sheet.append(["型号", "价格"])
    sheet.append(["A100", 1200])
    book.save(path)
    blocks = LoaderFactory.load(document(path))
    assert len(blocks) == 1
    assert blocks[0].source_location.cell_range == "A1:B2"
    assert blocks[0].source_location.sheet_name == "价格"
