from datetime import UTC, datetime
from pathlib import Path

from docx import Document
from docx.document import Document as WordFile
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


def new_document(title: str) -> WordFile:
    word = Document()
    section = word.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.2)
    for name in ["Normal", "Title", "Heading 1", "Heading 2"]:
        style = word.styles[name]
        style.font.name = "Microsoft YaHei"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    word.styles["Normal"].font.size = Pt(10)
    word.styles["Normal"].paragraph_format.space_after = Pt(6)
    word.styles["Title"].font.size = Pt(22)
    word.styles["Heading 1"].font.size = Pt(14)
    word.styles["Heading 2"].font.size = Pt(11)
    word.add_paragraph(title, "Title")
    word.add_paragraph("星屿示范企业 · 2026年9月版 · 全部内容为虚构演示数据")
    word.core_properties.author = "Traceable RAG 文档问答演示"
    word.core_properties.created = datetime(2026, 9, 1, tzinfo=UTC)
    word.core_properties.modified = datetime(2026, 9, 1, tzinfo=UTC)
    return word


def create_demo_data(output: Path) -> list[Path]:
    """生成原创虚构样例，不读取任何已有企业文件。"""
    output.mkdir(parents=True, exist_ok=True)
    policy = new_document("员工差旅管理制度")
    policy.add_paragraph("本制度说明星屿示范企业员工出差的费用标准和审批程序，供知识库演示使用。所有金额均为人民币。")
    policy.add_heading("第一章 适用范围", 1)
    policy.add_paragraph("适用于公司普通员工和管理人员在中国大陆地区的业务出差。")
    policy.add_heading("第二章 出差审批", 1)
    policy.add_paragraph("员工应提前3个工作日提交出差申请。预算不超过5000元由部门负责人审批，超过5000元还需总经理审批。")
    policy.add_heading("第三章 交通与餐补", 1)
    policy.add_paragraph("普通员工可乘坐高铁二等座或飞机经济舱。市内交通凭合规票据实报实销，每天最高100元。")
    policy.add_paragraph("餐补按实际出差天数发放，每人每天80元；客户已提供餐食的对应餐次不重复补贴。")
    policy.add_heading("第四章 住宿标准", 1)
    policy.add_heading("4 2 城市与职级标准", 2)
    policy.add_paragraph(
        "北京普通员工的住宿标准为600元每晚，上海普通员工同为600元每晚。管理人员在北京和上海的住宿标准为800元每晚。实际费用低于上限时按实报销。"
    )
    table = policy.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, ["城市", "职级", "住宿上限 元每晚"], strict=True):
        cell.text = value
    for row in [
        ["北京", "普通员工", "600"],
        ["上海", "普通员工", "600"],
        ["北京及上海", "管理人员", "800"],
        ["其他城市", "普通员工", "400"],
    ]:
        for cell, value in zip(table.add_row().cells, row, strict=True):
            cell.text = value
    for row in table.rows:
        props = row._tr.get_or_add_trPr()
        props.append(OxmlElement("w:cantSplit"))
    policy.add_heading("第五章 报销", 1)
    policy.add_paragraph("返程后10个工作日内提交发票与出差审批单。超出住宿标准的费用须事先获得总经理书面批准。")
    policy_path = output / "员工差旅管理制度.docx"
    policy.save(policy_path)

    contract = new_document("软件采购合同")
    contract.add_paragraph(
        "本虚构合同用于演示采购条款检索。甲方为星屿示范企业，乙方为云帆示范软件工作室，双方名称均为虚构。采购内容为20套文档管理软件许可，合同总价为120000元人民币。"
    )
    for heading, text in [
        ("第一条 付款方式", "合同生效后5个工作日内支付总价的30%，即36000元；软件验收合格后10个工作日内支付剩余70%，即84000元。"),
        ("第二条 交付与验收", "乙方应在收到首付款之日起20个工作日内完成软件交付和部署。甲方在交付后5个工作日内组织验收。"),
        (
            "第三条 延期交付责任",
            "因乙方原因延期交付，每逾期一天按合同总价的0.1%支付违约金，累计不超过合同总价的10%。逾期超过15个自然日，甲方有权解除合同并要求退还已付款。",
        ),
        (
            "第四条 售后服务",
            "自验收合格之日起提供12个月免费维护。工作日9点至18点提供支持，严重故障2小时内响应，24小时内提供解决方案或临时替代方案。",
        ),
        ("第五条 保密与范围", "双方仅在履行本合同所需范围内使用交付资料。本文件无真实签约主体，不构成可执行的商业合同。"),
    ]:
        contract.add_heading(heading, 1)
        contract.add_paragraph(text)
    contract_path = output / "软件采购合同.docx"
    contract.save(contract_path)

    book = Workbook()
    sheet = book.active
    sheet.title = "产品价格"
    sheet.append(["产品型号", "产品名称", "含税单价 元", "库存 件", "交期 工作日"])
    for row in [
        ["A100", "标准文档终端", 1299, 56, 3],
        ["A200", "专业文档终端", 2499, 12, 7],
        ["B100", "便携扫描配件", 699, 8, 5],
        ["C300", "会议显示组件", 3999, 35, 10],
    ]:
        sheet.append(row)
    table = Table(displayName="Products", ref="A1:E5")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium4", showRowStripes=True)
    sheet.add_table(table)
    sheet.freeze_panes = "A2"
    for column in ["A", "B", "C", "D", "E"]:
        sheet.column_dimensions[column].width = 24 if column == "B" else 21
    for row in sheet:
        sheet.row_dimensions[row[0].row].height = 29
        for cell in row:
            cell.font = Font(name="Microsoft YaHei", size=11, color="23362F")
            cell.alignment = Alignment(vertical="center")
            if cell.row == 1:
                cell.fill = PatternFill("solid", fgColor="315E42")
                cell.font = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
    for row in range(2, 6):
        sheet.cell(row, 3).number_format = '#,##0" 元"'
    sheet["A7"] = "说明：全部为虚构演示数据，价格和库存截至2026年9月1日。"
    sheet.merge_cells("A7:E7")
    sheet["A7"].font = Font(name="Microsoft YaHei", size=10, color="788570")
    book.properties.creator = "Traceable RAG 文档问答演示"
    # OOXML 工作簿时间戳不允许时区；明确从 UTC 转为无时区存储值。
    book.properties.created = datetime(2026, 9, 1, tzinfo=UTC).replace(tzinfo=None)
    price_path = output / "产品价格表.xlsx"
    book.save(price_path)
    return [policy_path, contract_path, price_path]


if __name__ == "__main__":
    for path in create_demo_data(Path(__file__).resolve().parents[1] / "data" / "demo"):
        print(path)
